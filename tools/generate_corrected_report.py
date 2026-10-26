from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_PATH = "docs/Informe_Del_comportamiento_a_la_implementacion_corregido.docx"


def set_cell_shading(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def set_cell_border(cell, color="D9D9D9"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), color)


def set_run_font(run, name="Aptos", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.add_run(text)
    return paragraph


def add_body(document, text, before=0, after=5):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.add_run(text)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)
    return paragraph


def add_code(document, code):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F5F7")
    set_cell_border(cell, "C9D2D9")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(code)
    set_run_font(run, name="Consolas", size=8.6)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, size=9, bold=True, color=(255, 255, 255))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F2F6FA")
            set_cell_border(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=9.2)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


document = Document()
section = document.sections[0]
section.top_margin = Cm(2.3)
section.bottom_margin = Cm(2.1)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)

normal = document.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
for name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
    style = document.styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Integración de la calidad en el PDA - Machly")
set_run_font(footer_run, size=8, color=(100, 100, 100))

for line, size, bold in [
    ("UNIVERSIDAD PRIVADA DEL VALLE", 13, True),
    ("Integración de la calidad en el PDA", 12, False),
]:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(line)
    set_run_font(run, name="Aptos Display", size=size, bold=bold)

document.add_paragraph().paragraph_format.space_after = Pt(17)
title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run("Del comportamiento a la implementación")
set_run_font(title_run, name="Aptos Display", size=23, bold=True)
subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(20)
subtitle_run = subtitle.add_run("Registro de usuario en Machly")
set_run_font(subtitle_run, size=13, italic=True)

cover = [
    ("Elaborado por", "Daniel Diaz Gorena\nJulio Guzman Almanza\nPablo Hugo Rocabdo"),
    ("Grupo", "B"),
    ("Fecha", "20 de octubre de 2026"),
    ("Asignatura", "Software Quality Assurance"),
]
table = document.add_table(rows=0, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for label, value in cover:
    cells = table.add_row().cells
    cells[0].text = label
    cells[1].text = value
    set_cell_shading(cells[0], "EAF1F7")
    for cell_index, cell in enumerate(cells):
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size=10, bold=(cell_index == 0))

document.add_page_break()

add_heading(document, "1. Objetivo")
add_body(
    document,
    "Aplicar ATDD, BDD y TDD en la funcionalidad de registro de usuario de Machly, "
    "integrando conceptos básicos de DDD. Machly es una aplicación Android desarrollada en Java; "
    "mantiene usuarios y roles en SQLite y dispone de una pantalla para registrar datos de usuario."
)

add_heading(document, "2. ATDD", 1)
by = document.add_paragraph()
by.add_run("Responsable: ").bold = True
by.add_run("Pablo Hugo Rocabdo")
add_body(
    document,
    "ATDD (Acceptance Test-Driven Development) es el desarrollo guiado por pruebas de aceptación. "
    "Antes de programar, cliente, desarrollo y pruebas acuerdan el comportamiento que debe entregar la funcionalidad. "
    "Ese acuerdo se expresa mediante criterios de aceptación observables."
)
add_body(document, "En esta actividad, Pablo Hugo Rocabdo asumió el rol de cliente. El registro de Machly se realiza desde la administración de usuarios.")
add_table(document, ["Criterio", "Historia de usuario", "Reglas de aceptación"], [
    ("CA-01 Registro exitoso", "Como administrador de Machly, quiero registrar un usuario con nombre, correo válido, contraseña de al menos 8 caracteres y rol, para que pueda disponer de una cuenta local.", "Nombre obligatorio; correo válido y no repetido; contraseña de 8 o más caracteres; usuario guardado al cumplir las reglas."),
    ("CA-02 Correo único", "Como administrador de Machly, quiero evitar el registro con un correo existente, para no tener dos cuentas con el mismo identificador.", "No se crea ni modifica una cuenta cuando el correo ya existe; se informa el motivo del rechazo."),
], widths=[3.1, 7.1, 6.0])
add_body(document, "Nota de calidad: estos criterios representan el comportamiento objetivo. Para que se cumplan desde la pantalla Android, el formulario debe invocar la validación de dominio antes de guardar.")

add_heading(document, "3. BDD", 1)
by = document.add_paragraph()
by.add_run("Responsable: ").bold = True
by.add_run("Julio Guzman Almanza")
add_body(
    document,
    "BDD (Behavior-Driven Development) traduce cada criterio de aceptación a ejemplos comprensibles para todo el equipo. "
    "Se emplea el formato Gherkin: Dado describe el contexto, Cuando la acción y Entonces el resultado esperado."
)
add_heading(document, "Escenario del criterio CA-01", 2)
add_code(document, '''Funcionalidad: Registro de usuario en Machly

Escenario: Registro exitoso con datos válidos
  Dado que no existe un usuario con el correo "ana@machly.com"
  Cuando registro a "Ana" con el correo "ana@machly.com"
    y la contraseña "abc12345"
  Entonces el sistema acepta el registro
    y guarda el usuario localmente
    y muestra el mensaje "User saved locally"''')
add_heading(document, "Escenario del criterio CA-02", 2)
add_code(document, '''Funcionalidad: Registro de usuario en Machly

Escenario: Intento de registro con correo duplicado
  Dado que ya existe un usuario con el correo "ana@machly.com"
  Cuando intento registrar un usuario con el mismo correo
    y una contraseña válida
  Entonces el sistema rechaza el registro
    y muestra el mensaje "El correo electrónico ya está registrado."''')

add_heading(document, "4. TDD", 1)
by = document.add_paragraph()
by.add_run("Responsable: ").bold = True
by.add_run("Daniel Diaz Gorena")
add_body(
    document,
    "TDD (Test-Driven Development) se aplicó al escenario de registro exitoso mediante el ciclo RED, GREEN y REFACTOR. "
    "La prueba se ejecuta como prueba unitaria JUnit y no depende de la interfaz Android."
)
add_heading(document, "Paso 1. RED - prueba que especifica el comportamiento", 2)
add_code(document, '''@Test
public void acceptsRegistrationWhenDataIsValidAndEmailIsAvailable() {
    UserRegistrationService service = new UserRegistrationService(email -> false);

    RegistrationValidationResult result = service.validateNewUser(
            "Ana", "ana@machly.com", "abc12345"
    );

    assertTrue(result.isValid());
}''')
add_body(document, "En la fase RED, esta prueba define el resultado esperado antes de implementar el servicio de registro. Sin UserRegistrationService ni RegistrationValidationResult, la prueba no podía compilar ni pasar.")
add_heading(document, "Paso 2. GREEN - implementación mínima", 2)
add_code(document, '''public RegistrationValidationResult validateNewUser(
        String name, String email, String password) {
    if (name == null || name.trim().isEmpty()) {
        return RegistrationValidationResult.invalid("El nombre es obligatorio.");
    }
    if (email == null || !EMAIL_PATTERN.matcher(email.trim()).matches()) {
        return RegistrationValidationResult.invalid("Ingresa un correo electrónico válido.");
    }
    if (password == null || password.length() < 8) {
        return RegistrationValidationResult.invalid("La contraseña debe tener al menos 8 caracteres.");
    }
    if (emailRegistrationChecker.isEmailRegistered(email.trim())) {
        return RegistrationValidationResult.invalid("El correo electrónico ya está registrado.");
    }
    return RegistrationValidationResult.valid();
}''')
add_body(document, "Con la implementación mínima, el registro válido es aceptado. La comprobación de correo duplicado se delega a EmailRegistrationChecker, que representa la consulta al repositorio.")
add_heading(document, "Paso 3. REFACTOR - separación de responsabilidades", 2)
add_body(document, "Se extrajo EmailRegistrationChecker para separar la regla de negocio de la tecnología de almacenamiento. La prueba de registro exitoso y una segunda prueba de correo repetido pasaron correctamente: 2 pruebas ejecutadas, 0 fallos y 0 errores.")
add_table(document, ["Prueba", "Resultado"], [
    ("acceptsRegistrationWhenDataIsValidAndEmailIsAvailable", "Aprobada"),
    ("rejectsRegistrationWhenEmailAlreadyExists", "Aprobada"),
], widths=[11.5, 4.7])

add_heading(document, "5. DDD", 1)
add_body(
    document,
    "DDD (Domain-Driven Design) permite que el equipo use las mismas palabras para conversar sobre el problema y para nombrar los elementos de código. "
    "El dominio tratado es el registro de un usuario en Machly."
)
add_table(document, ["Elemento", "Definición"], [
    ("Entidad", "Usuario: representa la cuenta registrada."),
    ("Atributos", "id, fullName, email, passwordHash, roleIdLocal y lastUpdated."),
    ("Servicio de dominio", "UserRegistrationService: valida la creación de un usuario."),
    ("Repositorio", "DbHelper: consulta y persiste usuarios en SQLite."),
    ("Reglas", "Nombre obligatorio; correo válido y único; contraseña de 8 o más caracteres."),
    ("Lenguaje ubicuo", "Usuario, Registro, Validación, Dominio y Repositorio."),
], widths=[4.0, 12.2])

add_heading(document, "6. Observaciones de calidad", 1)
add_bullet(document, "La aplicación conserva la contraseña como passwordHash; no la guarda directamente como texto plano.")
add_bullet(document, "La utilidad llamada BCryptUtils emplea actualmente SHA-256 y codificación Base64, no BCrypt. En un entorno productivo se recomienda BCrypt o Argon2 con sal y factor de trabajo.")
add_bullet(document, "La restricción UNIQUE en email de SQLite protege la integridad de datos, pero la regla de rechazo debe validarse antes de persistir para evitar actualizar un usuario existente durante un alta.")

add_heading(document, "Bibliografía", 1)
for reference in [
    "Ambler, S. W. (2012). Agile Modeling and Development. Ambysoft.",
    "Beck, K. (2003). Test Driven Development: By Example. Addison-Wesley.",
    "Evans, E. (2003). Domain-Driven Design: Tackling Complexity in the Heart of Software. Addison-Wesley.",
    "Freeman, S., & Pryce, N. (2009). Growing Object-Oriented Software, Guided by Tests. Addison-Wesley.",
    "North, D. (2006). Introducing BDD. Dan North & Associates.",
]:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(reference)

document.save(OUTPUT_PATH)
