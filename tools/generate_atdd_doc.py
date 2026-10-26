from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_PATH = "docs/ATDD_Parte_Individual.docx"


def set_cell_shading(cell, color):
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell_properties.append(shading)


def set_cell_border(cell, color="D9D9D9"):
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), color)


def format_cell(cell, bold=False, color=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(10)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor(*color)


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_heading(document, text):
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.add_run(text)
    return paragraph


document = Document()
section = document.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(11)
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.size = Pt(13)
styles["Heading 1"].font.bold = True
styles["Heading 1"].font.color.rgb = RGBColor(0, 0, 0)

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run("ATDD Registro de Usuario")
title_run.font.name = "Aptos Display"
title_run.font.size = Pt(22)
title_run.font.color.rgb = RGBColor(0, 0, 0)

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(14)
subtitle_run = subtitle.add_run("Parte individual - Desarrollo guiado por pruebas de aceptación")
subtitle_run.italic = True
subtitle_run.font.size = Pt(10.5)

intro = document.add_paragraph()
intro.paragraph_format.space_after = Pt(8)
intro.add_run("Propósito. ").bold = True
intro.add_run(
    "Definir, desde la perspectiva del cliente, el comportamiento aceptable para el registro de usuarios "
    "en la aplicación Machly. Esta sección aporta los criterios que el equipo utilizará posteriormente "
    "para redactar escenarios BDD e implementar pruebas TDD."
)

add_heading(document, "Concepto de ATDD")
paragraph = document.add_paragraph()
paragraph.paragraph_format.space_after = Pt(5)
paragraph.add_run("ATDD ").bold = True
paragraph.add_run(
    "(Acceptance Test Driven Development) es una práctica ágil en la que cliente, desarrollo y pruebas "
    "acuerdan los criterios de aceptación antes de construir una funcionalidad. Su objetivo es que el "
    "equipo comparta una definición clara de lo que debe hacer el sistema y pueda comprobarlo mediante pruebas."
)

add_heading(document, "Criterios de aceptación")

document.add_paragraph().add_run("CA 01 Registro exitoso").bold = True
criteria_one = document.add_paragraph()
criteria_one.paragraph_format.space_after = Pt(5)
criteria_one.add_run("Como ").bold = True
criteria_one.add_run("nuevo usuario, ")
criteria_one.add_run("quiero ").bold = True
criteria_one.add_run("registrarme con mi nombre, correo electrónico válido, contraseña de al menos 8 caracteres y un rol, ")
criteria_one.add_run("para ").bold = True
criteria_one.add_run("disponer de una cuenta almacenada en el sistema.")
for item in [
    "El nombre es obligatorio.",
    "El correo debe tener un formato válido.",
    "La contraseña debe contener al menos 8 caracteres.",
    "El correo no debe estar registrado previamente.",
    "Si todas las reglas se cumplen, el sistema guarda el usuario y confirma el registro.",
]:
    add_bullet(document, item)

document.add_paragraph().add_run("CA 02 Prevención de correo duplicado").bold = True
criteria_two = document.add_paragraph()
criteria_two.paragraph_format.space_after = Pt(5)
criteria_two.add_run("Como ").bold = True
criteria_two.add_run("administrador, ")
criteria_two.add_run("quiero ").bold = True
criteria_two.add_run("que el sistema rechace registros con un correo ya existente, ")
criteria_two.add_run("para ").bold = True
criteria_two.add_run("mantener la integridad de los datos y que cada usuario tenga un correo único.")
for item in [
    "Cuando un correo ya pertenece a otro usuario, el sistema no crea ni modifica registros.",
    "El sistema comunica al usuario que el correo electrónico ya está registrado.",
]:
    add_bullet(document, item)

add_heading(document, "Matriz de aceptación")
table = document.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["Criterio", "Condición a comprobar", "Resultado esperado"]
for index, label in enumerate(headers):
    cell = table.rows[0].cells[index]
    cell.text = label
    set_cell_shading(cell, "1F4E78")
    format_cell(cell, bold=True, color=(255, 255, 255))

rows = [
    ("CA 01", "Datos obligatorios y válidos", "El registro es aceptado y el usuario queda guardado."),
    ("CA 01", "Contraseña de 8 o más caracteres", "El registro puede continuar."),
    ("CA 02", "Correo ya registrado", "El registro es rechazado y se informa el motivo."),
]
for row_index, values in enumerate(rows):
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cells[index].text = value
        if row_index % 2 == 1:
            set_cell_shading(cells[index], "F2F6FA")
        format_cell(cells[index])

document.add_paragraph()
closing = document.add_paragraph()
closing.paragraph_format.space_before = Pt(5)
closing.add_run("Resultado esperado de esta parte. ").bold = True
closing.add_run(
    "Los criterios definidos sirven como acuerdo entre el rol cliente y el equipo. Los escenarios Gherkin "
    "y la implementación TDD corresponden a las demás partes del trabajo grupal."
)

document.save(OUTPUT_PATH)
